#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
神戸マツダ「ぐるっと点検ネット」計画書を Word(.docx) で生成するスクリプト。
日本語フォント(MS Pゴシック/游ゴシック相当)・見出しスタイル・表組み・
ブランドカラー(ソウルレッド)の装飾つき。
"""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SOUL = RGBColor(0xE1, 0x1D, 0x2A)      # ソウルレッド
INK = RGBColor(0x16, 0x18, 0x1C)       # ニアブラック
STEEL = RGBColor(0x5B, 0x62, 0x6B)     # グレー
JP_FONT = "游ゴシック"                  # 環境になければ Word 側で代替

doc = Document()

# ── 既定フォントを日本語対応に ──
style = doc.styles["Normal"]
style.font.name = JP_FONT
style.font.size = Pt(10.5)
style.element.rPr.rFonts.set(qn("w:eastAsia"), JP_FONT)


def set_jp(run, name=JP_FONT):
    run.font.name = name
    r = run._element
    r.rPr.rFonts.set(qn("w:eastAsia"), name)


def shade_cell(cell, hex_color):
    """セルに背景色を設定"""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def add_heading(text, level=1):
    p = doc.add_heading(level=level)
    run = p.add_run(text)
    set_jp(run)
    if level == 1:
        run.font.color.rgb = SOUL
        run.font.size = Pt(16)
    elif level == 2:
        run.font.color.rgb = INK
        run.font.size = Pt(13)
    else:
        run.font.color.rgb = STEEL
        run.font.size = Pt(11.5)
    return p


def add_para(text, bold=False, color=None, size=10.5, italic=False, align=None):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_jp(run)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    if align:
        p.alignment = align
    return p


def add_bullet(text, sub=False):
    p = doc.add_paragraph(style="List Bullet" if not sub else "List Bullet 2")
    run = p.add_run(text)
    set_jp(run)
    run.font.size = Pt(10.5)
    return p


def add_table(headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        run = hdr[i].paragraphs[0].add_run(h)
        set_jp(run)
        run.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade_cell(hdr[i], "16181C")
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(val))
            set_jp(run)
            run.font.size = Pt(9.5)
    if widths:
        for i, w in enumerate(widths):
            for row in table.rows:
                row.cells[i].width = Cm(w)
    return table


# ══════════════════════════════ 表紙 ══════════════════════════════
for _ in range(3):
    doc.add_paragraph()
add_para("神戸マツダ", bold=True, color=SOUL, size=20, align=WD_ALIGN_PARAGRAPH.CENTER)
add_para("「ぐるっと点検ネット」", bold=True, color=INK, size=26, align=WD_ALIGN_PARAGRAPH.CENTER)
add_para("企画・開発計画書", bold=True, color=INK, size=18, align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_paragraph()
add_para("複数店舗連携 車検・点検スケジューリング・プラットフォーム",
         color=STEEL, size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
for _ in range(6):
    doc.add_paragraph()
add_para("文書バージョン：1.0", color=STEEL, size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
add_para("作成日：2026-06-05", color=STEEL, size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
add_para("ステータス：コンセプト実証（プロトタイプ実装済み）",
         color=STEEL, size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_page_break()

# ══════════════════════════════ 概要表 ══════════════════════════════
add_table(
    ["項目", "内容"],
    [
        ["プロジェクト名", "ぐるっと点検ネット（複数店舗連携 車検・点検スケジューリング）"],
        ["対象企業", "神戸マツダ（兵庫県内 複数店舗展開）"],
        ["文書バージョン", "1.0"],
        ["作成日", "2026-06-05"],
        ["ステータス", "コンセプト実証（プロトタイプ実装済み）"],
    ],
    widths=[4, 12],
)

# ════════════════ 1. 背景と課題 ════════════════
add_heading("1. 背景と課題", 1)
add_heading("1.1 お客様側の課題", 2)
add_bullet("車検・点検は「めんどうなこと」という認識が根強い。")
add_bullet("「出かける前に済ませたい」という心理から、予約が午前中に集中する。")
add_bullet("結果として午後の入庫枠がぽっかり空く、需要の時間帯偏在が発生している。")
add_heading("1.2 店舗・運営側の課題", 2)
add_bullet("午前は満杯で予約が取れず、午後は人もピットも遊んでしまう（稼働ムラ）。")
add_bullet("複数店舗を持つ強みが、予約・人員運用に十分活かされていない。")
add_bullet("各店舗の入庫状況・スタッフ配置を横断的に把握する仕組みがない。")
add_heading("1.3 課題の定量化（プロトタイプ計測値）", 2)
add_bullet("ネットワーク全体の午前稼働率は約 80〜90% に対し、午後は約 40〜50%。")
add_bullet("午前・午後で 30〜40 ポイントの偏りが存在。")
add_bullet("この偏りの是正が、待ち時間短縮とリソース有効活用の鍵。")

# ════════════════ 2. コンセプト ════════════════
add_heading("2. コンセプト", 1)
add_para("「点検は “めんどう” から “おでかけのついで” へ。」", bold=True, color=SOUL, size=12)
add_para("兵庫県内の複数店舗をネットワークとして連携させ、お客様のお出かけ先の店舗で"
         "車検・点検を受けられる仕組みを構築する。")
add_bullet("姫路の人が神戸へ出かける日は、神戸の店舗で点検。")
add_bullet("神戸の人が姫路へ出かける日は、姫路の店舗で点検。")
add_bullet("午後の空き枠を可視化・訴求し、混雑のない時間帯へ自然に誘導する。")
add_bullet("これらのスケジュールを AI が横断的に管理・最適化する。")

# ════════════════ 3. 目的・ゴール（KPI） ════════════════
add_heading("3. 目的・ゴール（KPI）", 1)
add_table(
    ["区分", "指標", "目標イメージ"],
    [
        ["平準化", "午前・午後の稼働率の差", "30〜40pt → 15pt 以内"],
        ["稼働率", "午後の平均稼働率", "40〜50% → 65% 以上"],
        ["顧客体験", "平均待ち時間", "短縮（午後枠誘導により）"],
        ["店舗運営", "1台あたり人員稼働効率", "スタッフ応援配置で改善"],
        ["集客", "「ついで点検」での新規入庫", "他エリア顧客の取り込み"],
    ],
    widths=[3, 6, 7],
)

# ════════════════ 4. ターゲットユーザー ════════════════
add_heading("4. ターゲットユーザー", 1)
add_bullet("お客様（エンドユーザー）：お出かけ・買い物・観光のついでに点検を済ませたい層、待ち時間を減らしたい層。")
add_bullet("店舗スタッフ／本部（管理者）：入庫スケジュール・スタッフ配置を管理し、複数店舗のリソースを最適配分したい立場。")

# ════════════════ 5. 提供価値 ════════════════
add_heading("5. 提供価値（主要機能）", 1)
add_table(
    ["#", "機能", "対象", "概要"],
    [
        ["1", "お出かけ先で点検（店舗レコメンド）", "お客様", "目的地エリア・日付・点検内容から、近くて空いている店舗をAIがおすすめ順に提案"],
        ["2", "午後への誘導", "お客様", "午後の空き枠を強調し、待ち時間の少ない時間帯へ誘導"],
        ["3", "おでかけスポット情報", "お客様", "各店舗周辺のグルメ・観光・遊びを写真付きで楽しく紹介"],
        ["4", "午後予約クーポン", "お客様", "午後枠の予約で、提携スポット・自治体・商工会議所のクーポンを発行"],
        ["5", "電動自転車の無料貸出", "お客様", "点検中に周辺を回るおしゃれな電動アシスト自転車を無料レンタル"],
        ["6", "自治体・商工会議所連携（相互送客）", "地域", "まちと店舗でお客様を送りあい、地域全体を活性化"],
        ["7", "入庫スケジュール可視化", "管理者", "全店舗×日付の稼働をAM/PMヒートマップで把握"],
        ["8", "スタッフ店舗間配置", "管理者", "所属店舗を越えた応援シフトの管理とAI提案"],
        ["9", "AI最適化アシスタント", "管理者", "午後誘導・店舗振り分け・応援配置を重要度つきで自動提案"],
    ],
    widths=[1, 5, 2.5, 8],
)

# ════════════════ 5.1 午後予約特典・地域連携 ════════════════
add_heading("5.1 午後予約特典・地域連携（地域活性化）", 1)
add_para("午前偏重の是正と地域活性化を両立する、3つの「午後予約特典」。")
add_heading("① 午後予約クーポン", 2)
add_bullet("午後枠（13:00以降）の予約完了で、提携スポット・自治体・商工会議所と連携したクーポンを発行（例：明石焼き1皿サービス、姫路城入城料割引）。")
add_bullet("午後への自然な需要シフトを促すインセンティブとして機能。")
add_heading("② 電動自転車の無料貸し出し（「ぐるっとバイク」）", 2)
add_bullet("点検・車検でクルマを預けている間、おしゃれな電動アシスト自転車（Luupのようなスマートロック付き）を無料で貸し出し。各店に配備。")
add_bullet("クルマがなくても身軽に周辺スポットへ。坂の多い神戸や城下町の姫路でも快適。")
add_heading("③ 自治体・商工会議所との連携による相互送客", 2)
add_bullet("店舗 → まち：来店客に午後クーポンと電動自転車でまちへお出かけいただく。")
add_bullet("まち → 店舗：自治体・商工会議所の観光案内やイベントから「ついで点検」へ送客。")
add_bullet("人の流れと消費を生み、まちと店舗がともに潤う好循環＝地域活性化を目指す。")

# ════════════════ 6. 画面構成 ════════════════
add_heading("6. 画面構成（サイトマップ）", 1)
add_para("お客様向けと管理者向けで、レイアウト・ナビゲーションを完全分離。")
add_heading("6.1 お客様サイト（明るい写真ドリブンUI）", 2)
add_table(
    ["URL", "画面", "内容"],
    [
        ["/", "ホーム", "フォトヒーロー、エリアタイル、カテゴリタイル、課題と解決、店舗一覧"],
        ["/availability", "空き店舗をさがす", "エリア・日付・点検内容で検索→AIおすすめ順に店舗提示"],
        ["/stores", "店舗一覧", "写真カードで全店舗とおでかけスポット"],
        ["/stores/[id]", "店舗詳細", "フォトヒーロー、周辺スポット、空き状況、午後クーポン、電動自転車、スタッフ体制"],
        ["/perks", "おでかけ特典・地域連携", "午後クーポン一覧、電動自転車レンタル、自治体・商工会議所連携"],
    ],
    widths=[4, 4, 8],
)
add_heading("6.2 スタッフ管理ポータル（ダークな業務用UI）", 2)
add_table(
    ["URL", "画面", "内容"],
    [
        ["/staff", "ダッシュボード", "稼働サマリー、AM/PMヒートマップ、本日の店舗別状況"],
        ["/staff/members", "スタッフ配置", "所属/応援可能、AI応援提案、スタッフ名簿"],
        ["/staff/ai", "AI最適化", "提案フィルタ、ネットワーク診断"],
    ],
    widths=[4, 4, 8],
)
add_para("※ 相互リンク（「スタッフの方→」「←お客様ページ」）で行き来可能。", color=STEEL, size=9.5)

# ════════════════ 7. AI最適化エンジン ════════════════
add_heading("7. AI最適化エンジン", 1)
add_para("ルールベースの最適化ロジックを src/lib/optimizer.ts に実装。")
add_heading("7.1 店舗レコメンド（お客様向け）", 2)
add_para("目的地との近さ・午後の空き・全体の空き・スタッフ充足度をスコアリング（0–100）し、"
         "おすすめ順に提示。午後枠を加点して午前偏重を是正。")
add_heading("7.2 午後誘導の提案", 2)
add_bullet("検出条件：午前稼働 ≥ 85% かつ 午後稼働 ≤ 45%。")
add_bullet("提案：午前希望客へ「午後来店で待ち時間ほぼゼロ＋特典」を案内。")
add_heading("7.3 近隣店舗への振り分け", 2)
add_bullet("検出条件：逼迫店舗（稼働 ≥ 85%）の 25km 圏内に余裕（稼働 < 60%）のある店舗。")
add_bullet("提案：お出かけ動線上の顧客を空き店舗へ案内。")
add_heading("7.4 スタッフ応援配置", 2)
add_bullet("検出条件：入庫予定に対して人員不足の店舗（1人あたり処理目安 6台）。")
add_bullet("提案：手すきの近隣店舗から、応援可能なスタッフを配置。")
add_para("将来的には、予約実績・天候・キャンペーン情報を学習し、LLM（Claude等）による"
         "自然言語提案・自動配信へ拡張する。", italic=True, color=STEEL, size=10)

# ════════════════ 8. システム構成 ════════════════
add_heading("8. システム構成・技術スタック", 1)
add_table(
    ["区分", "採用技術"],
    [
        ["フレームワーク", "Next.js 16（App Router）"],
        ["言語", "TypeScript"],
        ["UI", "React 19 / Tailwind CSS v4"],
        ["画像", "next/image（Unsplash を remotePatterns 配信）"],
        ["データ", "src/lib/ の決定論的シード（DB不要・再現性あり）"],
        ["ホスティング", "Vercel を想定"],
    ],
    widths=[4, 12],
)
add_heading("ディレクトリ構成", 3)
tree = (
    "src/\n"
    "├── app/\n"
    "│   ├── (customer)/   お客様サイト（独立レイアウト）\n"
    "│   └── (staff)/      スタッフ管理ポータル（独立レイアウト）\n"
    "├── components/       共通UI（ヘッダー・カード・PhotoCard 等）\n"
    "└── lib/\n"
    "    ├── types.ts      型定義\n"
    "    ├── stores.ts     店舗マスタ・写真・距離計算\n"
    "    ├── staff.ts      スタッフマスタ（店舗間応援関係）\n"
    "    ├── schedule.ts   入庫枠・稼働の生成（午前偏重を再現）\n"
    "    └── optimizer.ts  AI最適化エンジン"
)
pcode = doc.add_paragraph()
rcode = pcode.add_run(tree)
rcode.font.name = "Consolas"
rcode.font.size = Pt(9)

# ════════════════ 9. データモデル ════════════════
add_heading("9. データモデル（主要エンティティ）", 1)
add_bullet("Store：店舗（エリア・座標・ピット数・テーマ・写真・スポット）")
add_bullet("Staff：スタッフ（所属店舗・応援可能店舗・資格・スキル）")
add_bullet("Slot / DayLoad：時間帯ごとの入庫枠と稼働集計")
add_bullet("Suggestion：AI提案（種別・重要度・対象店舗・効果）")
add_bullet("Availability：お客様向けの空き状況（おすすめ度・理由つき）")

# ════════════════ 10. デザイン方針 ════════════════
add_heading("10. デザイン方針", 1)
add_bullet("ブランド基調：マツダらしいブラック（#16181C）＆シルバー。差し色にソウルレッド（#E11D2A）。")
add_bullet("兵庫の風景カラー：海（青）・山（緑）・港/城（アンバー）でエリアを彩る。")
add_bullet("トーン：旅メディア（tabiiro.jp）を参考に、写真ドリブンで明るく楽しい雰囲気。")
add_bullet("役割分担：赤＝行動（予約系CTA）、黒＝構造・ブランド、シルバー＝面・境界。")

# ════════════════ 11. ロードマップ ════════════════
add_heading("11. 想定スケジュール（ロードマップ）", 1)
add_table(
    ["フェーズ", "内容", "状態"],
    [
        ["Phase 0", "コンセプト実証プロトタイプ（本リポジトリ）", "完了"],
        ["Phase 1", "実データ連携（店舗・スタッフ・予約DB）、認証", "計画"],
        ["Phase 2", "実予約システム／基幹システム連携", "計画"],
        ["Phase 3", "AIのLLM化・需要予測の学習導入", "計画"],
        ["Phase 4", "顧客アプリ／通知・キャンペーン自動配信", "構想"],
    ],
    widths=[3, 10, 3],
)

# ════════════════ 12. 拡張候補 ════════════════
add_heading("12. 今後の拡張候補", 1)
add_bullet("スタッフ管理ポータルへのログイン認証（顧客と管理画面の分離強化）。")
add_bullet("予約実績・天候・イベント情報を取り込んだ需要予測。")
add_bullet("Claude API 連携による自然言語でのAI提案・要約・自動メッセージ生成。")
add_bullet("地図UI・経路連携（お出かけ動線上の店舗提案）。")
add_bullet("キャンペーン（午後来店特典・洗車サービス等）の自動最適化。")

# ════════════════ 13. リスク・留意点 ════════════════
add_heading("13. リスク・留意点", 1)
add_bullet("本リポジトリはコンセプト実証用デモ。店舗・スタッフ・スポット・写真はサンプルを含む。")
add_bullet("データは固定基準日（2026-06-05）の決定論的シードで、実予約とは連動しない。")
add_bullet("実運用には基幹システム・予約DB・認証・個人情報保護対応が別途必要。")
add_bullet("外部写真（Unsplash）は実運用時に権利確認の上、自社撮影素材等への差し替えを推奨。")

# ════════════════ 14. デプロイ ════════════════
add_heading("14. デプロイ", 1)
add_bullet("Vercel に GitHub 連携で Import（追加設定ほぼ不要、環境変数も現状不要）。")
add_bullet("本番ブランチの設定、または main へのマージで公開。")
add_bullet("フィーチャーブランチへの push ごとにプレビューURLが自動生成される。")

doc.add_paragraph()
add_para("本計画書は「ぐるっと点検ネット」コンセプト実証フェーズの記録です。",
         italic=True, color=STEEL, size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER)

OUT = "docs/神戸マツダ_ぐるっと点検ネット_計画書.docx"
doc.save(OUT)
print("saved:", OUT)
